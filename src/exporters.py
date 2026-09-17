"""
exporters.py — Export OCR results into MANY formats the user can choose.

Formats: txt, md, json, html, csv, xlsx, docx

DOCX is written WITHOUT any external library: a .docx is just a ZIP of XML, so
this module builds the minimal Office Open XML itself. That means DOCX ALWAYS
works even if `python-docx` is not installed on the server. (If python-docx IS
installed, it is used for nicer table formatting; otherwise we fall back to the
built-in writer.)
"""
from __future__ import annotations

import json
import logging
import os
import tempfile
import zipfile
from pathlib import Path
from typing import Optional
from xml.sax.saxutils import escape

import pandas as pd


LOG = logging.getLogger("ocr.exporters")

ALL_FORMATS = [
    "txt",
    "md",
    "json",
    "html",
    "csv",
    "xlsx",
    "docx",
]

TABLE_ONLY_FORMATS = {
    "csv",
    "xlsx",
}

def _atomic_write_text(
    destination: Path,
    content: str,
    encoding: str = "utf-8",
) -> Path:
    """
    Write text to a temporary file and replace the destination only
    after the write completes successfully.
    """

    destination.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.stem}_",
        suffix=".tmp",
        dir=str(destination.parent),
    )

    temporary_path = Path(temporary_name)

    try:
        with os.fdopen(
            file_descriptor,
            "w",
            encoding=encoding,
            newline="",
        ) as file_handle:
            file_handle.write(content)
            file_handle.flush()
            os.fsync(file_handle.fileno())

        temporary_path.replace(destination)

        return destination

    except Exception:
        temporary_path.unlink(
            missing_ok=True,
        )
        raise


def _atomic_output_path(
    destination: Path,
) -> tuple[Path, Path]:
    """
    Return a temporary output path and its final destination.

    This helper is used by libraries that require a file path,
    such as pandas and python-docx.
    """

    destination.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    temporary_path = (
        destination.parent
        / (
            f".{destination.stem}_"
            f"{next(tempfile._get_candidate_names())}"
            f"{destination.suffix}"
        )
    )

    return temporary_path, destination


def _replace_output(
    temporary_path: Path,
    destination: Path,
) -> Path:
    """
    Replace the destination only after file generation succeeds.
    """

    if not temporary_path.is_file():
        raise FileNotFoundError(
            f"Expected export file was not created: "
            f"{temporary_path.name}"
        )

    if temporary_path.stat().st_size == 0:
        temporary_path.unlink(
            missing_ok=True,
        )

        raise ValueError(
            f"Generated export file is empty: "
            f"{destination.name}"
        )

    temporary_path.replace(destination)

    return destination

def _markdown_cell(value) -> str:
    """
    Convert a value into a safe Markdown table cell.
    """

    if pd.isna(value):
        return ""

    return (
        str(value)
        .replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\r", " ")
        .replace("\n", " ")
    )


def _mdtable(
    dataframe: pd.DataFrame,
) -> str:
    columns = [
        _markdown_cell(column)
        for column in dataframe.columns
    ]

    lines = [
        "| " + " | ".join(columns) + " |",
        "| "
        + " | ".join(
            ["---"] * len(columns)
        )
        + " |",
    ]

    for _, row in dataframe.iterrows():
        values = [
            _markdown_cell(value)
            for value in row.tolist()
        ]

        lines.append(
            "| " + " | ".join(values) + " |"
        )

    return "\n".join(lines)


def export_all(
    outdir,
    stem: str,
    text: str = "",
    table_df: Optional[pd.DataFrame] = None,
    summary: Optional[dict] = None,
    formats: Optional[list[str]] = None,
    locale: str = "auto",
) -> dict[str, str]:
    """
    Export OCR output into the requested formats.

    Returns:
        Dictionary mapping format names to successfully generated
        file paths.

    Notes:
        CSV and XLSX require a non-empty table.
        Failures are logged and do not prevent other formats from
        being generated.
    """

    del locale

    output_directory = Path(outdir)
    output_directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    safe_stem = (
        Path(stem).name.strip("._")
        or "ocr_result"
    )

    requested_formats = (
        formats or ["txt"]
    )

    normalized_formats = []

    for file_format in requested_formats:
        normalized = str(
            file_format
        ).strip().lower()

        if (
            normalized in ALL_FORMATS
            and normalized
            not in normalized_formats
        ):
            normalized_formats.append(
                normalized
            )

    has_table = (
        table_df is not None
        and not table_df.empty
    )

    written: dict[str, str] = {}

    for file_format in normalized_formats:
        if (
            file_format in TABLE_ONLY_FORMATS
            and not has_table
        ):
            LOG.info(
                "Skipping table-only export "
                "format=%s reason=no_table",
                file_format,
            )
            continue

        try:
            if file_format == "txt":
                destination = (
                    output_directory
                    / f"{safe_stem}.txt"
                )

                _atomic_write_text(
                    destination,
                    text or "",
                )

            elif file_format == "md":
                destination = (
                    output_directory
                    / f"{safe_stem}.md"
                )

                sections = []

                if text:
                    sections.append(text)

                if has_table:
                    sections.append(
                        _mdtable(table_df)
                    )

                body = "\n\n".join(sections)

                _atomic_write_text(
                    destination,
                    body,
                )

            elif file_format == "json":
                destination = (
                    output_directory
                    / f"{safe_stem}.json"
                )

                payload = {
                    "text": text or "",
                    "table": (
                        [
                            [
                                str(column)
                                for column
                                in table_df.columns
                            ]
                        ]
                        + table_df.where(
                            pd.notna(table_df),
                            None,
                        ).values.tolist()
                        if has_table
                        else None
                    ),
                    "summary": summary or {},
                }

                json_content = json.dumps(
                    payload,
                    indent=2,
                    ensure_ascii=False,
                    default=str,
                )

                _atomic_write_text(
                    destination,
                    json_content,
                )

            elif file_format == "html":
                destination = (
                    output_directory
                    / f"{safe_stem}.html"
                )

                parts = [
                    "<!DOCTYPE html>",
                    '<html lang="en">',
                    "<head>",
                    '<meta charset="utf-8">',
                    (
                        "<meta name=\"viewport\" "
                        "content=\"width=device-width, "
                        "initial-scale=1\">"
                    ),
                    "<title>OCR Result</title>",
                    "</head>",
                    "<body>",
                ]

                if text:
                    parts.append(
                        "<pre>"
                        + _html_escape(text)
                        + "</pre>"
                    )

                if has_table:
                    parts.append(
                        table_df.to_html(
                            index=False,
                            border=1,
                            escape=True,
                            na_rep="",
                        )
                    )

                parts.extend([
                    "</body>",
                    "</html>",
                ])

                _atomic_write_text(
                    destination,
                    "\n".join(parts),
                )

            elif file_format == "csv":
                destination = (
                    output_directory
                    / f"{safe_stem}.csv"
                )

                temporary_path, final_path = (
                    _atomic_output_path(
                        destination
                    )
                )

                try:
                    table_df.to_csv(
                        temporary_path,
                        index=False,
                        encoding="utf-8-sig",
                    )

                    _replace_output(
                        temporary_path,
                        final_path,
                    )

                except Exception:
                    temporary_path.unlink(
                        missing_ok=True,
                    )
                    raise

            elif file_format == "xlsx":
                destination = (
                    output_directory
                    / f"{safe_stem}_tables.xlsx"
                )

                temporary_path, final_path = (
                    _atomic_output_path(
                        destination
                    )
                )

                try:
                    with pd.ExcelWriter(
                        temporary_path,
                        engine="openpyxl",
                    ) as excel_writer:
                        table_df.to_excel(
                            excel_writer,
                            sheet_name="Table_1",
                            index=False,
                        )

                        if summary:
                            summary_frame = (
                                pd.DataFrame(
                                    [
                                        {
                                            key: _summary_value(
                                                value
                                            )
                                            for key, value
                                            in summary.items()
                                        }
                                    ]
                                )
                            )

                            summary_frame.to_excel(
                                excel_writer,
                                sheet_name="_Summary",
                                index=False,
                            )

                    _replace_output(
                        temporary_path,
                        final_path,
                    )

                except Exception:
                    temporary_path.unlink(
                        missing_ok=True,
                    )
                    raise

            elif file_format == "docx":
                destination = (
                    output_directory
                    / f"{safe_stem}.docx"
                )

                temporary_path, final_path = (
                    _atomic_output_path(
                        destination
                    )
                )

                try:
                    _write_docx(
                        temporary_path,
                        text,
                        table_df,
                        has_table,
                    )

                    _replace_output(
                        temporary_path,
                        final_path,
                    )

                except Exception:
                    temporary_path.unlink(
                        missing_ok=True,
                    )
                    raise

            else:
                continue

            if not destination.is_file():
                raise FileNotFoundError(
                    "The exporter did not create "
                    f"{destination.name}."
                )

            written[file_format] = str(
                destination
            )

            LOG.info(
                "Export completed "
                "format=%s filename=%s "
                "size_bytes=%s",
                file_format,
                destination.name,
                destination.stat().st_size,
            )

        except Exception:
            LOG.exception(
                "Export failed "
                "format=%s stem=%s",
                file_format,
                safe_stem,
            )

    return written

def _summary_value(value):
    """
    Convert nested summary values into spreadsheet-friendly text.
    """

    if isinstance(
        value,
        (dict, list, tuple, set),
    ):
        return json.dumps(
            value,
            ensure_ascii=False,
            default=str,
        )

    return value


def _html_escape(value: str) -> str:
    return (
        str(value)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&#x27;")
    )

# ---------------------------------------------------------------------------
# DOCX writer
# ---------------------------------------------------------------------------
def _write_docx(path: Path, text: str, df, has: bool) -> Path:
    """Prefer python-docx (nicer tables); otherwise build a valid .docx by hand."""
    try:
        from docx import Document
        doc = Document()
        if text:
            for line in (text.splitlines() or [""]):
                doc.add_paragraph(line)
        if has:
            rows, cols = df.shape
            t = doc.add_table(rows=rows + 1, cols=cols)
            t.style = "Table Grid"
            for j, c in enumerate(df.columns):
                t.cell(0, j).text = str(c)
            for i in range(rows):
                for j in range(cols):
                    t.cell(i + 1, j).text = str(df.iat[i, j])
        doc.save(path)
        return path
    except Exception:
        # Fallback: no python-docx (or it failed) -> build minimal OOXML zip.
        return _write_docx_raw(path, text, df, has)


def _p(text: str) -> str:
    """
    Build one Word paragraph.
    """

    escaped_text = escape(
        str(text)
    )

    return (
        "<w:p>"
        "<w:r>"
        "<w:t xml:space='preserve'>"
        + escaped_text
        + "</w:t>"
        "</w:r>"
        "</w:p>"
    )


def _table_xml(df) -> str:
    """A simple bordered Word table."""
    borders = ("<w:tblBorders>"
               + "".join(f"<w:{s} w:val='single' w:sz='4' w:space='0' w:color='000000'/>"
                         for s in ("top", "left", "bottom", "right",
                                   "insideH", "insideV"))
               + "</w:tblBorders>")
    tblpr = f"<w:tblPr><w:tblW w:w='0' w:type='auto'/>{borders}</w:tblPr>"

    def cell(value) -> str:
        if pd.isna(value):
            cell_text = ""
        else:
            cell_text = str(value)

        escaped_text = escape(cell_text)

        return (
            "<w:tc>"
            "<w:tcPr>"
            "<w:tcW w:w='0' w:type='auto'/>"
            "</w:tcPr>"
            "<w:p>"
            "<w:r>"
            "<w:t xml:space='preserve'>"
            + escaped_text
            + "</w:t>"
            "</w:r>"
            "</w:p>"
            "</w:tc>"
        )


    rows = ["<w:tr>" + "".join(cell(c) for c in df.columns) + "</w:tr>"]
    for _, r in df.iterrows():
        rows.append("<w:tr>" + "".join(cell(v) for v in r.tolist()) + "</w:tr>")
    return f"<w:tbl>{tblpr}{''.join(rows)}</w:tbl>"


def _write_docx_raw(path: Path, text: str, df, has: bool) -> Path:
    body = []
    for line in (text.splitlines() if text else [""]):
        body.append(_p(line))
    if has:
        body.append(_table_xml(df))
    body.append("<w:sectPr/>")

    document_xml = (
        "<?xml version='1.0' encoding='UTF-8' standalone='yes'?>"
        "<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
        f"<w:body>{''.join(body)}</w:body></w:document>"
    )
    content_types = (
        "<?xml version='1.0' encoding='UTF-8' standalone='yes'?>"
        "<Types xmlns='http://schemas.openxmlformats.org/package/2006/content-types'>"
        "<Default Extension='rels' ContentType='application/vnd.openxmlformats-package.relationships+xml'/>"
        "<Default Extension='xml' ContentType='application/xml'/>"
        "<Override PartName='/word/document.xml' "
        "ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml'/>"
        "</Types>"
    )
    rels = (
        "<?xml version='1.0' encoding='UTF-8' standalone='yes'?>"
        "<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>"
        "<Relationship Id='rId1' "
        "Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument' "
        "Target='word/document.xml'/></Relationships>"
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document_xml)
    return path



